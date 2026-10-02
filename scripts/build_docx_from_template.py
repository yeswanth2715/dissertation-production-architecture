from __future__ import annotations

import argparse
import shutil
from pathlib import Path

from common import latest_file, load_manifest, project_path, read_text, rel


def resolve_template(project_dir: Path, explicit: str | None) -> Path:
    if explicit:
        path = Path(explicit)
        return path if path.is_absolute() else project_dir / path

    manifest = load_manifest(project_dir)
    configured = manifest.get("template", {}).get("source") if isinstance(manifest.get("template"), dict) else None
    if configured:
        path = Path(str(configured))
        if not path.is_absolute():
            path = project_dir / path
        if path.exists():
            return path

    local = latest_file(project_dir, ["*template*.docx", "*TEMPLATE*.docx"])
    if local:
        return local

    repo = latest_file(project_dir.parents[1] / "templates" / "source", ["*.docx"])
    if repo:
        return repo

    raise FileNotFoundError("No DOCX template found. Provide --template or set template.source in project.yaml.")


def append_markdown(docx_path: Path, markdown_path: Path) -> None:
    try:
        from docx import Document
    except Exception as exc:
        raise RuntimeError(f"python-docx unavailable: {exc}") from exc

    doc = Document(docx_path)
    for raw in read_text(markdown_path).splitlines():
        line = raw.strip()
        if not line:
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
        else:
            doc.add_paragraph(line)
    doc.save(docx_path)


def build(project: str | Path, template: str | None, output: str | None, body_md: str | None, force: bool) -> Path:
    root = project_path(project)
    template_path = resolve_template(root, template)
    out = Path(output) if output else root / "outputs" / "dissertation.docx"
    if not out.is_absolute():
        out = root / out
    out.parent.mkdir(parents=True, exist_ok=True)
    if out.exists() and not force:
        raise FileExistsError(f"Output already exists: {out}. Use --force to overwrite.")

    shutil.copy2(template_path, out)
    if body_md:
        body_path = Path(body_md)
        if not body_path.is_absolute():
            body_path = root / body_path
        append_markdown(out, body_path)
    return out


def main() -> None:
    parser = argparse.ArgumentParser(description="Duplicate a dissertation DOCX template and optionally append project-specific Markdown.")
    parser.add_argument("--project", required=True)
    parser.add_argument("--template")
    parser.add_argument("--output")
    parser.add_argument("--body-md", help="Optional project-specific Markdown body to append. The script does not generate prose.")
    parser.add_argument("--force", action="store_true")
    args = parser.parse_args()

    out = build(args.project, args.template, args.output, args.body_md, args.force)
    print(f"created {rel(out)}")


if __name__ == "__main__":
    main()
