from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from common import dump_json, latest_file, load_manifest, project_path, rel, word_count, write_text


DOC_EXTS = {".doc", ".docx", ".pdf", ".md", ".txt"}
DATA_EXTS = {".xlsx", ".xls", ".csv", ".tsv", ".json"}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".svg"}


def docx_summary(path: Path) -> dict[str, Any]:
    try:
        from docx import Document
    except Exception as exc:
        return {"path": rel(path), "error": f"python-docx unavailable: {exc}"}

    try:
        doc = Document(path)
    except Exception as exc:
        return {"path": rel(path), "error": f"cannot read docx: {exc}"}

    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    headings = [
        p.text.strip()
        for p in doc.paragraphs
        if p.text.strip() and p.style is not None and p.style.name.startswith("Heading")
    ]
    text = "\n".join(paragraphs)
    upper = text.upper()
    return {
        "path": rel(path),
        "paragraphs": len(paragraphs),
        "word_count_estimate": word_count(text),
        "tables": len(doc.tables),
        "headings": headings[:80],
        "has_introduction": "INTRODUCTION" in upper,
        "has_chapter_one": "CHAPTER ONE" in upper,
        "has_chapter_two": "CHAPTER TWO" in upper,
        "has_chapter_three": "CHAPTER THREE" in upper or "METHODOLOGY" in upper,
        "has_chapter_four": "CHAPTER FOUR" in upper or "FINDINGS" in upper,
        "has_conclusion": "CONCLUDING REMARKS" in upper or "CONCLUSION" in upper,
        "has_bibliography": "BIBLIOGRAPHY" in upper or "REFERENCES" in upper,
    }


def classify_method(manifest: dict[str, Any], files: list[Path]) -> str:
    configured = manifest.get("method", {})
    if isinstance(configured, dict) and configured.get("type"):
        return str(configured["type"])
    text = " ".join(p.name.lower() for p in files)
    if any(token in text for token in ["survey", "questionnaire", "responses"]):
        return "survey-only"
    if any(p.suffix.lower() in {".xlsx", ".csv", ".tsv"} for p in files):
        return "data-driven"
    return "unknown"


def scan(project: str | Path) -> dict[str, Any]:
    root = project_path(project)
    manifest = load_manifest(root)
    files = [p for p in root.rglob("*") if p.is_file() and not p.name.startswith("~$")]
    by_ext = Counter(p.suffix.lower() or "[none]" for p in files)
    by_area: dict[str, list[str]] = defaultdict(list)
    for path in files:
        try:
            parts = path.relative_to(root).parts
            top = parts[0] if len(parts) > 1 else "."
        except Exception:
            top = "."
        by_area[top].append(rel(path))

    docx_files = [p for p in files if p.suffix.lower() == ".docx"]
    summaries = [docx_summary(p) for p in sorted(docx_files, key=lambda f: rel(f))]
    method = classify_method(manifest, files)

    latest_docx = latest_file(root / "outputs" if (root / "outputs").exists() else root, ["*.docx"])
    latest_results = latest_file(root, ["*.xlsx", "*.csv", "*.tsv"])

    return {
        "project": root.name,
        "project_path": rel(root),
        "manifest_present": bool(manifest),
        "configured_method": manifest.get("method", {}).get("type") if isinstance(manifest.get("method"), dict) else None,
        "detected_method": method,
        "file_count": len(files),
        "extensions": dict(sorted(by_ext.items())),
        "areas": {k: sorted(v) for k, v in sorted(by_area.items())},
        "latest_output_docx": rel(latest_docx) if latest_docx else None,
        "latest_result_file": rel(latest_results) if latest_results else None,
        "docx_summaries": summaries,
    }


def discovery_markdown(result: dict[str, Any]) -> str:
    docx_lines = []
    for item in result["docx_summaries"]:
        flags = []
        for key, label in [
            ("has_introduction", "Introduction"),
            ("has_chapter_one", "Chapter One"),
            ("has_chapter_two", "Chapter Two"),
            ("has_chapter_three", "Chapter Three/Methodology"),
            ("has_chapter_four", "Chapter Four/Findings"),
            ("has_conclusion", "Conclusion"),
            ("has_bibliography", "Bibliography"),
        ]:
            if item.get(key):
                flags.append(label)
        docx_lines.append(
            f"- `{item['path']}`: {item.get('word_count_estimate', 'n/a')} words est.; "
            f"{item.get('tables', 'n/a')} tables; detected sections: {', '.join(flags) if flags else 'none'}."
        )

    area_lines = []
    for area, paths in result["areas"].items():
        area_lines.append(f"- `{area}/`: {len(paths)} file(s)")

    return (
        "# Project Discovery\n\n"
        f"Project: `{result['project']}`\n\n"
        f"Detected method: `{result['detected_method']}`\n\n"
        f"Manifest present: `{result['manifest_present']}`\n\n"
        f"Latest output DOCX: `{result['latest_output_docx']}`\n\n"
        f"Latest result file: `{result['latest_result_file']}`\n\n"
        "## File Areas\n\n"
        + "\n".join(area_lines)
        + "\n\n## DOCX Readiness Signals\n\n"
        + ("\n".join(docx_lines) if docx_lines else "- No DOCX files detected.")
        + "\n\n## Machine Summary\n\n"
        "```json\n"
        + dump_json(
            {
                "file_count": result["file_count"],
                "extensions": result["extensions"],
                "latest_output_docx": result["latest_output_docx"],
                "latest_result_file": result["latest_result_file"],
            }
        )
        + "\n```\n"
    )


def main() -> None:
    parser = argparse.ArgumentParser(description="Scan a dissertation project and produce a compact discovery summary.")
    parser.add_argument("--project", required=True, help="Project folder name under projects/ or a direct path.")
    parser.add_argument("--json", action="store_true", help="Print JSON instead of Markdown.")
    parser.add_argument("--write", action="store_true", help="Write project-discovery.md in the project folder.")
    args = parser.parse_args()

    result = scan(args.project)
    if args.write:
        write_text(project_path(args.project) / "project-discovery.md", discovery_markdown(result))

    print(dump_json(result) if args.json else discovery_markdown(result))


if __name__ == "__main__":
    main()
