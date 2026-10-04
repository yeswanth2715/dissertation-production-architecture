from __future__ import annotations

import argparse
import re
import sys
import zipfile
from pathlib import Path
from typing import Any

from common import artifact_path, dissertation_path, dump_json, latest_file, load_manifest, project_path, read_text, rel, word_count, write_text
from utils.rules import load_banned_terms


FORBIDDEN_FINAL_TERMS = [
    "pre-survey",
    "when responses are supplied",
    "project 1",
    "project 2",
    "project 3",
    "project 4",
    "project 5",
    "project 6",
    "project 7",
    "/learn",
    "artifact folder",
    "fallback",
    "response-quality",
    "completion timestamp",
    "straightline",
    "low variation",
    "survey was weak",
    "ai-generated",
    "decorative artwork",
    "time taken",
    "duration",
]
FORBIDDEN_FINAL_TERMS = sorted(set(FORBIDDEN_FINAL_TERMS + load_banned_terms()))

PLACEHOLDER_PATTERNS = [
    r"\[[^\]]*(insert|add|todo|placeholder|personalised|personalized)[^\]]*\]",
    r"\bTBD\b",
    r"\bTODO\b",
    r"\bXXX\b",
]


def result(status: str, name: str, detail: str) -> dict[str, str]:
    return {"status": status, "check": name, "detail": detail}


def load_docx(path: Path):
    try:
        from docx import Document
    except Exception as exc:
        raise RuntimeError(f"python-docx unavailable: {exc}") from exc
    return Document(path)


def docx_text(doc: Any) -> str:
    texts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                texts.append(cell.text)
    return "\n".join(texts)


def heading_texts(doc: Any) -> list[str]:
    return [
        p.text.strip()
        for p in doc.paragraphs
        if p.text.strip() and p.style is not None and p.style.name.startswith("Heading")
    ]


def body_word_count(doc: Any) -> int:
    start = None
    end = None
    for index, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip().upper()
        style = paragraph.style.name if paragraph.style is not None else ""
        if start is None and style.startswith("Heading") and text == "INTRODUCTION":
            start = index
        if start is not None and style.startswith("Heading") and text in {"BIBLIOGRAPHY", "REFERENCES"}:
            end = index
            break
    if start is None:
        start = 0
    if end is None:
        end = len(doc.paragraphs)
    total = sum(word_count(p.text) for p in doc.paragraphs[start:end])
    total += sum(word_count(" ".join(cell.text for row in table.rows for cell in row.cells)) for table in doc.tables)
    return total


def media_count(path: Path) -> int:
    try:
        with zipfile.ZipFile(path) as archive:
            return len([n for n in archive.namelist() if n.startswith("word/media/")])
    except Exception:
        return 0


def zip_ok(path: Path) -> bool:
    try:
        with zipfile.ZipFile(path) as archive:
            return archive.testzip() is None
    except Exception:
        return False


def find_docx(project_dir: Path, explicit: str | None) -> Path:
    return dissertation_path(project_dir, explicit)


def survey_result_files(project_dir: Path) -> list[Path]:
    survey_terms = ("survey", "questionnaire", "response", "responses", "respondent", "participant")
    ignored_parts = {"outputs", "tables", "figures", "logs", "code"}
    candidates = []
    for path in project_dir.rglob("*"):
        if not path.is_file() or path.suffix.lower() not in {".xlsx", ".xls", ".csv", ".tsv"}:
            continue
        rel_parts = {part.lower() for part in path.relative_to(project_dir).parts[:-1]}
        name = path.name.lower()
        if rel_parts.intersection(ignored_parts) and not any(term in name for term in survey_terms):
            continue
        if any(term in name for term in survey_terms):
            candidates.append(path)
    return sorted(candidates)


def expected_word_range(manifest: dict[str, Any]) -> tuple[int | None, int | None, int | None]:
    wc = manifest.get("word_count", {}) if isinstance(manifest.get("word_count"), dict) else {}
    target = wc.get("target")
    if not isinstance(target, int):
        return None, None, None
    min_pct = float(wc.get("pre_humaniser_min_pct", 0.85))
    max_pct = float(wc.get("pre_humaniser_max_pct", 0.92))
    return target, int(target * min_pct), int(target * max_pct)


def qa(project: str | Path, docx: str | None = None, stage: str = "draft") -> dict[str, Any]:
    root = project_path(project)
    manifest = load_manifest(root)
    checks: list[dict[str, str]] = []

    for folder in ["inputs", "knowledge-base", "feedback", "qa", "outputs"]:
        path = root / folder
        checks.append(result("pass" if path.exists() else "fail", f"folder:{folder}", rel(path)))

    checks.append(result("pass" if (root / "project.yaml").exists() else "warn", "project manifest", "project.yaml present" if (root / "project.yaml").exists() else "project.yaml missing"))
    checks.append(result("pass" if (root / "project-discovery.md").exists() else "warn", "project discovery", "project-discovery.md present" if (root / "project-discovery.md").exists() else "run scripts/scan_project.py --write"))

    docx_path = find_docx(root, docx)
    doc_summary: dict[str, Any] = {"path": None}
    if not docx_path.is_file():
        checks.append(result("fail" if stage == "final" else "warn", "output docx", "No output DOCX found."))
    else:
        checks.append(result("pass" if zip_ok(docx_path) else "fail", "docx package", rel(docx_path)))
        doc = load_docx(docx_path)
        text = docx_text(doc)
        lower = text.lower()
        headings = heading_texts(doc)
        body_words = body_word_count(doc)
        target, min_words, max_words = expected_word_range(manifest)
        if target:
            detail = f"{body_words} words; target {target}; pre-humaniser range {min_words}-{max_words}"
            status = "fail" if body_words > target else ("pass" if min_words <= body_words <= max_words else "warn")
        else:
            detail = f"{body_words} words; no manifest target"
            status = "warn"
        checks.append(result(status, "body word count", detail))
        checks.append(result("pass" if len(doc.tables) > 0 else "warn", "tables", f"{len(doc.tables)} table(s)"))
        checks.append(result("pass" if media_count(docx_path) > 0 else "warn", "figures/media", f"{media_count(docx_path)} embedded media file(s)"))

        required = ["INTRODUCTION", "CHAPTER ONE", "CHAPTER TWO", "CHAPTER THREE"]
        if stage == "final":
            required.extend(["CHAPTER FOUR"])
        for item in required:
            checks.append(result("pass" if item in text.upper() else "fail", f"required section:{item}", item))
        has_conclusion = "CONCLUSION" in text.upper() or "CONCLUDING REMARKS" in text.upper()
        checks.append(result("pass" if has_conclusion else "fail", "required section:CONCLUSION", "CONCLUSION"))

        if "CHAPTER FOUR" in text.upper():
            strict_ch4 = ["4.1 Findings", "4.2 Analysis", "4.3 Discussion"]
            empirical_ch4 = [
                "4.1 Experimental Overview",
                "4.2 Overall Detection Performance",
                "4.3 Per-Class Performance",
                "4.4 Confidence-Threshold Analysis",
                "4.5 Accuracy-Latency Trade-Off",
                "4.6 Model Comparison",
                "4.7 Error Analysis",
                "4.8 Robustness Analysis",
                "4.9 Object-Size Analysis",
                "4.10 Decision-State Evaluation",
                "4.11 Statistical Analysis",
                "4.12 Discussion Against the Literature",
            ]
            if all(item.lower() in lower for item in empirical_ch4):
                checks.append(result("pass", "chapter four empirical structure", "expanded empirical findings/analysis/discussion structure present"))
            else:
                for item in strict_ch4:
                    checks.append(result("pass" if item.lower() in lower else "fail", f"chapter four heading:{item}", item))
        bad_terms = [term for term in FORBIDDEN_FINAL_TERMS if term in lower]
        checks.append(result("pass" if not bad_terms else "fail", "professor-facing internal wording", ", ".join(bad_terms) if bad_terms else "no blocked terms found"))

        placeholders = []
        for pattern in PLACEHOLDER_PATTERNS:
            placeholders.extend(re.findall(pattern, text, flags=re.IGNORECASE))
        checks.append(result("pass" if not placeholders else ("fail" if stage == "final" else "warn"), "placeholder scan", "none" if not placeholders else f"{len(placeholders)} possible placeholder(s)"))

        citations = re.findall(r"\([A-Z][A-Za-z'.-]+(?:\s+and\s+[A-Z][A-Za-z'.-]+| et al\.)?,\s*(?:19|20)\d{2}\)", text)
        has_bib = "BIBLIOGRAPHY" in text.upper() or "REFERENCES" in text.upper()
        checks.append(result("pass" if citations and has_bib else "warn", "citation presence", f"{len(citations)} in-text citation pattern(s); bibliography={has_bib}"))

        doc_summary = {
            "path": rel(docx_path),
            "body_words": body_words,
            "tables": len(doc.tables),
            "media_files": media_count(docx_path),
            "heading_count": len(headings),
            "heading_preview": headings[:30],
        }

    outputs = root / "outputs"
    configured_results = manifest.get("survey", {}).get("results")
    survey_results = []
    if configured_results:
        configured_path = artifact_path(root, str(configured_results))
        if configured_path.is_file():
            survey_results = [configured_path]
    has_results = bool(survey_results)
    method = manifest.get("method", {}).get("type") if isinstance(manifest.get("method"), dict) else "unknown"
    if method in {"survey", "survey-only", "mixed-methods"} or has_results:
        detail = ", ".join(rel(path) for path in survey_results) if has_results else "not found"
        checks.append(result("pass" if has_results else "warn", "survey result file", detail))
        if stage == "final" and not has_results:
            checks.append(result("fail", "survey final gate", "final survey-based findings need returned results"))

    failed = [c for c in checks if c["status"] == "fail"]
    warned = [c for c in checks if c["status"] == "warn"]
    return {
        "project": root.name,
        "stage": stage,
        "docx": doc_summary,
        "summary": {"pass": len([c for c in checks if c["status"] == "pass"]), "warn": len(warned), "fail": len(failed)},
        "checks": checks,
    }


def markdown(report: dict[str, Any]) -> str:
    lines = [
        "# Fast QA Report",
        "",
        f"Project: `{report['project']}`",
        f"Stage: `{report['stage']}`",
        "",
        "## Summary",
        "",
        f"- Pass: {report['summary']['pass']}",
        f"- Warn: {report['summary']['warn']}",
        f"- Fail: {report['summary']['fail']}",
        "",
        "## Document",
        "",
        "```json",
        str(report["docx"]).replace("'", '"'),
        "```",
        "",
        "## Checks",
        "",
        "| Status | Check | Detail |",
        "| --- | --- | --- |",
    ]
    for check in report["checks"]:
        detail = check["detail"].replace("|", "\\|").replace("\n", " ")
        lines.append(f"| {check['status']} | {check['check']} | {detail} |")
    lines.append("")
    return "\n".join(lines)


def main() -> None:
    parser = argparse.ArgumentParser(description="Run deterministic Fast QA for a dissertation project.")
    parser.add_argument("--project", required=True)
    parser.add_argument("--docx")
    parser.add_argument("--stage", default="draft", choices=["draft", "pre-survey", "final"])
    parser.add_argument("--json", action="store_true")
    parser.add_argument("--write", action="store_true", help="Write qa/fast-qa-report.md.")
    parser.add_argument("--strict", action="store_true", help="Exit non-zero when failures are found.")
    args = parser.parse_args()

    report = qa(args.project, args.docx, args.stage)
    if args.write:
        write_text(project_path(args.project) / "qa" / "fast-qa-report.md", markdown(report))
    print(dump_json(report) if args.json else markdown(report))
    if args.strict and report["summary"]["fail"]:
        sys.exit(1)


if __name__ == "__main__":
    main()
