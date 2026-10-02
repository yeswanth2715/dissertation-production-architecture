from __future__ import annotations

import json
import zipfile
from pathlib import Path
from typing import Any

from common import rel, word_count
from utils.rules import load_json_config


TEXT_PREVIEW_LIMIT = 1200
IGNORED_PARTS = {"__pycache__", ".git"}


def _read_text_preview(path: Path) -> dict[str, Any]:
    try:
        text = path.read_text(encoding="utf-8", errors="replace")
    except Exception as exc:
        return {"read_status": "warn", "detail": f"text read failed: {exc}"}
    return {
        "read_status": "pass",
        "word_count": word_count(text),
        "preview": text[:TEXT_PREVIEW_LIMIT],
    }


def _inspect_docx(path: Path) -> dict[str, Any]:
    try:
        from docx import Document
    except Exception as exc:
        return {"read_status": "warn", "detail": f"python-docx unavailable: {exc}"}
    try:
        doc = Document(path)
    except Exception as exc:
        return {"read_status": "warn", "detail": f"docx read failed: {exc}"}
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    headings = [
        p.text.strip()
        for p in doc.paragraphs
        if p.text.strip() and p.style is not None and p.style.name.startswith("Heading")
    ]
    return {
        "read_status": "pass",
        "paragraphs": len(paragraphs),
        "word_count_estimate": word_count("\n".join(paragraphs)),
        "tables": len(doc.tables),
        "headings_preview": headings[:30],
    }


def _inspect_xlsx(path: Path) -> dict[str, Any]:
    try:
        import openpyxl  # type: ignore
    except Exception:
        try:
            with zipfile.ZipFile(path) as archive:
                worksheets = [name for name in archive.namelist() if name.startswith("xl/worksheets/")]
            return {"read_status": "pass", "package": "xlsx", "worksheet_count": len(worksheets)}
        except Exception as exc:
            return {"read_status": "warn", "detail": f"xlsx package read failed: {exc}"}
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        sheets = [{"name": ws.title, "max_row": ws.max_row, "max_column": ws.max_column} for ws in workbook.worksheets]
        workbook.close()
        return {"read_status": "pass", "sheets": sheets}
    except Exception as exc:
        return {"read_status": "warn", "detail": f"xlsx read failed: {exc}"}


def inspect_artifact(path: Path, project_root: Path) -> dict[str, Any]:
    suffix = path.suffix.lower()
    stat = path.stat()
    relative = path.relative_to(project_root)
    top_area = relative.parts[0] if len(relative.parts) > 1 else "."
    result: dict[str, Any] = {
        "path": rel(path),
        "area": top_area,
        "extension": suffix or "[none]",
        "size_bytes": stat.st_size,
    }
    config = load_json_config("artifact_requirements.json", {}).get("final_output_gate", {})
    text_exts = set(config.get("readable_text_extensions", []))
    docx_exts = set(config.get("inspectable_document_extensions", []))
    data_exts = set(config.get("inspectable_data_extensions", []))

    if suffix in text_exts:
        result.update(_read_text_preview(path))
    elif suffix in docx_exts:
        result.update(_inspect_docx(path))
    elif suffix in data_exts:
        result.update(_inspect_xlsx(path))
    else:
        result.update({"read_status": "inventory-only", "detail": "binary or unsupported artifact inventoried"})
    return result


def collect_artifacts(project_root: Path) -> list[dict[str, Any]]:
    artifacts: list[dict[str, Any]] = []
    for path in sorted(project_root.rglob("*")):
        if not path.is_file() or path.name.startswith("~$"):
            continue
        if any(part in IGNORED_PARTS for part in path.parts):
            continue
        artifacts.append(inspect_artifact(path, project_root))
    return artifacts


def gate(project_root: Path, stage: str = "final") -> dict[str, Any]:
    artifacts = collect_artifacts(project_root)
    areas = sorted({item["area"] for item in artifacts})
    config = load_json_config("artifact_requirements.json", {}).get("final_output_gate", {})
    required_areas = config.get("required_existing_areas_to_scan", [])
    checks: list[dict[str, str]] = []
    checks.append({"status": "pass" if artifacts else "fail", "check": "artifact inventory", "detail": f"{len(artifacts)} artifact(s) inspected"})
    for area in required_areas:
        status = "pass" if area in areas else "warn"
        checks.append({"status": status, "check": f"area scanned:{area}", "detail": "present" if area in areas else "not present in project"})
    unreadable = [item for item in artifacts if item.get("read_status") == "warn"]
    checks.append({"status": "pass" if not unreadable else "warn", "check": "artifact readability", "detail": f"{len(unreadable)} warning(s)"})
    failures = [item for item in checks if item["status"] == "fail"]
    warnings = [item for item in checks if item["status"] == "warn"]
    return {
        "project": project_root.name,
        "stage": stage,
        "status": "fail" if failures else ("warn" if warnings else "pass"),
        "summary": {"artifacts": len(artifacts), "areas": areas, "warnings": len(warnings), "failures": len(failures)},
        "checks": checks,
        "artifacts": artifacts,
    }


def markdown(report: dict[str, Any]) -> str:
    lines = [
        "# Artifact Gate Report",
        "",
        f"Project: `{report['project']}`",
        f"Stage: `{report['stage']}`",
        f"Status: `{report['status']}`",
        "",
        "This report records the architecture requirement to inspect existing project artifacts before final output generation.",
        "",
        "## Checks",
        "",
        "| Status | Check | Detail |",
        "| --- | --- | --- |",
    ]
    for check in report["checks"]:
        lines.append(f"| {check['status']} | {check['check']} | {check['detail'].replace('|', '\\|')} |")
    lines.extend(["", "## Artifacts", "", "| Area | Status | Path | Detail |", "| --- | --- | --- | --- |"])
    for item in report["artifacts"]:
        detail = item.get("detail") or f"{item.get('word_count', item.get('word_count_estimate', ''))} words"
        lines.append(f"| {item['area']} | {item.get('read_status')} | `{item['path']}` | {str(detail).replace('|', '\\|')} |")
    return "\n".join(lines) + "\n"


def to_json(report: dict[str, Any]) -> str:
    return json.dumps(report, indent=2, ensure_ascii=True, default=str)
